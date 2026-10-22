# -*- coding: utf-8 -*-
"""Exportação do relatório para Excel (.xlsx) ou CSV."""
import csv
import os


def _fmt_data(d):
    return d.strftime("%d/%m/%Y") if d else ""


def cabecalho_resumo(res):
    cab = ["Usuário"]
    if len(res["camadas"]) > 1:
        cab += list(res["camadas"])
    cab += ["Total", "Dias com produção", "Média/dia", "% do total", "Primeiro registro", "Último registro"]
    if res["sem_data"]:
        cab.append("Sem data")
    return cab


def linhas_resumo(res):
    linhas = []
    for r in res["resumo"]:
        lin = [r["usuario"]]
        if len(res["camadas"]) > 1:
            lin += r["por_camada"]
        lin += [r["total"], r["dias"], round(r["media"], 1), round(r["pct"], 1),
                _fmt_data(r["primeiro"]), _fmt_data(r["ultimo"])]
        if res["sem_data"]:
            lin.append(r["sem_data"])
        linhas.append(lin)
    tot = ["TOTAL"]
    if len(res["camadas"]) > 1:
        tot += [sum(r["por_camada"][i] for r in res["resumo"]) for i in range(len(res["camadas"]))]
    tot += [res["total"], res["dias_com_producao"], round(res["media_geral_dia"], 1), 100.0, "", ""]
    if res["sem_data"]:
        tot.append(res["sem_data"])
    linhas.append(tot)
    return linhas


def cabecalho_periodo(res):
    return ["Usuário"] + res["rotulos"] + ["Total"]


def linhas_periodo(res):
    linhas = [[m["usuario"]] + m["valores"] + [m["total"]] for m in res["matriz"]]
    linhas.append(["TOTAL"] + res["totais_periodo"] + [sum(res["totais_periodo"])])
    return linhas


def exportar_xlsx(caminho, res, info):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    azul = PatternFill("solid", fgColor="1F4E78")
    cinza = PatternFill("solid", fgColor="D9E1F2")
    branco = Font(bold=True, color="FFFFFF")
    negrito = Font(bold=True)
    fino = Side(style="thin", color="A6A6A6")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)

    def aba(ws, titulo, cab, linhas):
        ws.append([titulo])
        ws["A1"].font = Font(bold=True, size=13)
        for txt in info:
            ws.append([txt])
        ws.append([])
        lin0 = ws.max_row + 1
        ws.append(cab)
        for c in range(1, len(cab) + 1):
            cel = ws.cell(row=lin0, column=c)
            cel.fill, cel.font, cel.border = azul, branco, borda
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for lin in linhas:
            ws.append(lin)
            for c in range(1, len(cab) + 1):
                ws.cell(row=ws.max_row, column=c).border = borda
        for c in range(1, len(cab) + 1):
            cel = ws.cell(row=ws.max_row, column=c)
            cel.fill, cel.font = cinza, negrito
        for c in range(1, len(cab) + 1):
            larg = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(lin0, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, larg + 2), 40)
        ws.freeze_panes = ws.cell(row=lin0 + 1, column=2)

    ws1 = wb.active
    ws1.title = "Resumo por usuário"
    aba(ws1, "Relatório de produção — resumo por usuário", cabecalho_resumo(res), linhas_resumo(res))

    ws2 = wb.create_sheet("Por período")
    aba(ws2, "Relatório de produção — usuário x período", cabecalho_periodo(res), linhas_periodo(res))

    wb.save(caminho)
    return [caminho]


def exportar_csv(caminho, res, info):
    base, _ = os.path.splitext(caminho)
    arqs = [(base + "_resumo.csv", cabecalho_resumo(res), linhas_resumo(res)),
            (base + "_periodo.csv", cabecalho_periodo(res), linhas_periodo(res))]
    for arq, cab, linhas in arqs:
        with open(arq, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            for txt in info:
                w.writerow([txt])
            w.writerow([])
            w.writerow(cab)
            for lin in linhas:
                w.writerow([str(v).replace(".", ",") if isinstance(v, float) else v for v in lin])
    return [a[0] for a in arqs]
